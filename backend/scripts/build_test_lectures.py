"""Builds two more sample lectures for the AI evaluation (Phase 8, AI 1, #73).

Not part of the app. The built files are committed in tests/samples, so this
only needs running again to change them. Needs LibreOffice for the PDF.
Run from backend/:

    python -m scripts.build_test_lectures

The text is adapted, mostly word for word, from Dive into Deep Learning
(Zhang, Lipton, Li and Smola, https://d2l.ai), chapters 3.1 and 7.1 to 7.5,
which is licensed CC BY-SA 4.0; both lectures say so on their last page and
carry the same licence. They add a second topic and two layouts to the test
set: a slide deck with a title on every slide, and lecture notes as a
multi-page PDF, the format lecturers upload most.
"""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from app.services.image_captioning import find_libreoffice

SAMPLES = Path(__file__).resolve().parent.parent / "tests" / "samples"
CREDIT = (
    "Adapted from Dive into Deep Learning by Aston Zhang, Zachary C. Lipton, Mu Li and "
    "Alexander J. Smola (https://d2l.ai), licensed CC BY-SA 4.0. This adaptation is "
    "shared under the same licence."
)

CNN_SLIDES: list[tuple[str, list[str]]] = [
    (
        "Convolutional Neural Networks: Convolutions, Padding, Stride and Pooling",
        ["Deep Learning, Week 7", "From fully connected layers to convolutions"],
    ),
    (
        "Why not a fully connected layer for images?",
        [
            "A one-megapixel photograph gives each input one million dimensions.",
            "Even an aggressive reduction to one thousand hidden dimensions needs a fully "
            "connected layer with 10^6 x 10^3 = 10^9 parameters.",
            "Learning that many parameters would need many GPUs and an enormous dataset.",
            "Images have rich structure, and CNNs exploit that structure to learn with "
            "fewer parameters.",
        ],
    ),
    (
        "Invariance and locality",
        [
            "What Waldo looks like does not depend on where Waldo is located.",
            "Translation invariance: the earliest layers respond similarly to the same "
            "patch, regardless of where it appears in the image.",
            "Locality: the earliest layers focus on local regions, without regard for the "
            "contents of distant regions.",
            "Deeper layers capture longer-range features, like higher level vision in nature.",
        ],
    ),
    (
        "The cross-correlation operation",
        [
            "Convolutional layers are a misnomer: the operation they express is a "
            "cross-correlation.",
            "The kernel window starts at the upper-left corner of the input and slides "
            "left to right and top to bottom.",
            "At each position the input window and the kernel are multiplied elementwise "
            "and summed into a single scalar.",
            "Example: a 3 x 3 input and a 2 x 2 kernel give a 2 x 2 output; the first "
            "element is 0x0 + 1x1 + 3x2 + 4x3 = 19.",
        ],
    ),
    (
        "Output size without padding",
        [
            "An input of shape nh x nw and a kernel of shape kh x kw give an output of "
            "shape (nh - kh + 1) x (nw - kw + 1).",
            "The output is smaller because the kernel must fit wholly within the image.",
            "A convolutional layer cross-correlates the input and kernel and adds a scalar "
            "bias; its two parameters are the kernel and the bias.",
            "Kernels are usually initialised randomly, as in a fully connected layer.",
        ],
    ),
    (
        "Edge detection with a kernel",
        [
            "Image: 6 x 8 pixels, the middle four columns black (0) and the rest white (1).",
            "Kernel K = [1, -1], of height 1 and width 2, is a finite difference operator.",
            "The output is 1 for an edge from white to black, -1 for an edge from black to "
            "white, and 0 everywhere else.",
            "Applied to the transposed image the output vanishes: K only detects vertical edges.",
            "For larger kernels and deeper networks, the kernel is learned from data "
            "rather than designed by hand.",
        ],
    ),
    (
        "Padding",
        [
            "Successive convolutions shrink the image: ten layers of 5 x 5 convolutions "
            "reduce a 240 x 240 image to 200 x 200, slicing off 30% of it.",
            "Pixels in the corners are hardly used at all.",
            "Padding adds extra pixels of filler, usually zeros, around the boundary.",
            "With ph rows and pw columns of padding the output is "
            "(nh - kh + ph + 1) x (nw - kw + pw + 1).",
            "Setting ph = kh - 1 and pw = kw - 1 keeps the output the same size as the input.",
        ],
    ),
    (
        "Why odd kernel sizes?",
        [
            "CNNs commonly use kernels with odd height and width, such as 1, 3, 5 or 7.",
            "With an odd kernel the same number of rows can be padded on top and bottom, "
            "and the same number of columns on left and right.",
            "The output Y[i, j] is then computed with the window centred on X[i, j].",
            "Example: a 3 x 3 kernel with 1 pixel of padding on all sides keeps an 8 x 8 "
            "input at 8 x 8.",
        ],
    ),
    (
        "Stride",
        [
            "Stride is the number of rows and columns the window moves per slide.",
            "So far the stride has been 1 for both height and width.",
            "A larger stride skips intermediate locations, for computational efficiency or "
            "to downsample.",
            "With stride sh and sw the output shape is "
            "floor((nh - kh + ph + sh) / sh) x floor((nw - kw + pw + sw) / sw).",
        ],
    ),
    (
        "Pooling layers",
        [
            "Pooling mitigates the sensitivity of convolutional layers to location and "
            "spatially downsamples representations.",
            "A pooling window slides over the input like a kernel, but the pooling layer "
            "contains no parameters.",
            "Maximum pooling takes the largest value in the window; average pooling takes "
            "the mean.",
            "Example: 2 x 2 max-pooling of a 3 x 3 input gives max(0, 1, 3, 4) = 4 as its "
            "first output.",
            "In almost all cases, max-pooling is preferable to average pooling.",
        ],
    ),
    (
        "Pooling: padding, stride and channels",
        [
            "Like convolutional layers, pooling layers change the output shape, and padding "
            "and stride adjust it.",
            "Deep learning frameworks default to a pooling stride equal to the window size: "
            "a 3 x 3 window gets a stride of 3.",
            "With several input channels, pooling is applied to each channel separately, so "
            "the number of output channels equals the number of input channels.",
        ],
    ),
    ("Sources", [CREDIT]),
]

NOTES: list[tuple[int, str, list[str]]] = [
    (1, "Linear Regression", ["Lecture notes, Machine Learning, Week 2"]),
    (
        2,
        "Regression problems",
        [
            "Regression problems pop up whenever we want to predict a numerical value. "
            "Common examples include predicting prices of homes or stocks, predicting the "
            "length of stay of patients in hospital, and forecasting demand for retail sales.",
            "As a running example, suppose we wish to estimate the prices of houses in "
            "dollars from their area in square feet and their age in years. The dataset is "
            "called a training set, and each row, one sale, is called an example. The "
            "thing we are trying to predict, the price, is called the label or target. The "
            "variables the prediction is based on, area and age, are called features.",
        ],
    ),
    (
        2,
        "Assumptions",
        [
            "Linear regression is both the simplest and the most popular of the standard "
            "tools for regression, dating back to the start of the 19th century. It assumes "
            "first that the relationship between the features x and the target y is "
            "approximately linear: the conditional mean E[Y | X = x] is a weighted sum of "
            "the features. The target may still deviate from its expected value because of "
            "observation noise.",
            "Second, it assumes that the noise is well behaved, following a Gaussian distribution.",
        ],
    ),
    (
        2,
        "The model",
        [
            "The model predicts price = w_area * area + w_age * age + b. Here w_area and "
            "w_age are called weights and b is called the bias, offset or intercept. The "
            "weights determine the influence of each feature on the prediction. The bias "
            "is the value of the estimate when all features are zero; it is needed so the "
            "model can express every linear function of the features, not only lines that "
            "pass through the origin.",
            "Strictly speaking the model is an affine transformation: a linear "
            "transformation of the features by a weighted sum, combined with a translation "
            "by the bias. With d features collected in a vector x, the prediction is the "
            "dot product y_hat = w^T x + b. For a whole dataset, the design matrix X has one "
            "row for every example and one column for every feature, and the predictions "
            "are y_hat = Xw + b.",
        ],
    ),
    (
        2,
        "The squared error loss",
        [
            "Loss functions quantify the distance between the real and predicted values of "
            "the target. The loss is usually nonnegative, smaller is better, and a perfect "
            "prediction has a loss of 0. For regression the most common loss is the squared "
            "error, l = 1/2 (y_hat - y)^2. The constant 1/2 makes no real difference but "
            "cancels out when the loss is differentiated.",
            "Because of its quadratic form, large differences between estimates and "
            "targets lead to even larger contributions to the loss. This encourages the "
            "model to avoid large errors, but also makes it excessively sensitive to "
            "anomalous data. Over the whole training set the losses are averaged, and "
            "training seeks the weights and bias that minimise that average.",
        ],
    ),
    (
        2,
        "The analytic solution",
        [
            "Unlike most models, linear regression can be solved analytically. Folding the "
            "bias into the weights by appending a column of 1s to the design matrix, the "
            "problem is to minimise ||y - Xw||^2. The optimal weights are "
            "w* = (X^T X)^-1 X^T y.",
            "This solution is unique only when the matrix X^T X is invertible, that is, "
            "when the columns of the design matrix are linearly independent and no feature "
            "is linearly dependent on the others. Analytic solutions are rare: requiring "
            "one would exclude almost all of deep learning.",
        ],
    ),
    (
        2,
        "Minibatch stochastic gradient descent",
        [
            "Gradient descent reduces the error iteratively, updating the parameters in "
            "the direction that lowers the loss. Computing the gradient over the whole "
            "dataset before every update is extremely slow. Stochastic gradient descent "
            "(SGD) updates after a single example, but processors multiply and add numbers "
            "much faster than they move data from memory, and some layers, such as batch "
            "normalisation, need more than one example at a time.",
            "The solution is a minibatch: a small random sample of examples. A minibatch "
            "size between 32 and 256, preferably a multiple of a large power of 2, is a "
            "good start. In each iteration a minibatch is sampled, the gradient of its "
            "average loss is computed, multiplied by a small positive learning rate, and "
            "subtracted from the parameters.",
            "The minibatch size and the learning rate are hyperparameters: they are set by "
            "the user rather than updated in the training loop, and can be tuned, for "
            "example by Bayesian optimisation. The quality of the result is assessed on a "
            "separate validation set.",
        ],
    ),
    (
        2,
        "Vectorisation for speed",
        [
            "When training, a whole minibatch of examples is processed at once. Adding two "
            "vectors of 10,000 elements with a Python for-loop, one element at a time, is "
            "far slower than one vectorised addition, which is orders of magnitude faster. "
            "Vectorised code is also less error prone, because it is shorter and leaves "
            "the arithmetic to the library.",
        ],
    ),
    (
        2,
        "The normal distribution and squared loss",
        [
            "The squared loss can be motivated by assuming the observations come from noisy "
            "measurements, y = w^T x + b + noise, where the noise follows a normal "
            "distribution with mean 0 and standard deviation sigma.",
            "Maximising the likelihood of the data under this assumption is the same as "
            "minimising the negative log-likelihood, and that is the same as minimising "
            "the mean squared error. So under Gaussian noise, minimising the squared error "
            "is the maximum likelihood estimator of a linear model.",
        ],
    ),
    (
        2,
        "Linear regression as a neural network",
        [
            "Linear regression can be drawn as a single-layer neural network. The inputs "
            "are the features x1 to xd, and there is one output, o1. Every input is "
            "connected to the output, so the layer is a fully connected layer, also called "
            "a dense layer. The number of layers counts only the layers that compute, so "
            "the network has one layer.",
            "A biological neuron receives signals from other neurons through its "
            "dendrites, weights them by synaptic strengths, sums them in the cell body, and "
            "sends the result along its axon. Modern deep learning draws more on "
            "mathematics, statistics and computer science than on neuroscience.",
        ],
    ),
    (
        2,
        "Summary",
        [
            "Linear regression fits a weighted sum of the features plus a bias by "
            "minimising the squared error, usually with minibatch stochastic gradient "
            "descent. Its key ingredients, a model, a loss function and an optimisation "
            "algorithm, are the ingredients of nearly every deep learning model.",
        ],
    ),
    (2, "Sources", [CREDIT]),
]


def build_cnn_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    deck = Presentation()
    title_layout, content_layout = deck.slide_layouts[0], deck.slide_layouts[1]
    for number, (title, bullets) in enumerate(CNN_SLIDES):
        slide = deck.slides.add_slide(title_layout if number == 0 else content_layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
        for paragraph in body.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18 if number else 24)
    deck.save(path)


def build_notes_pdf(path: Path) -> None:
    """The notes as a Word document with real headings, one section to a page,
    saved as a PDF by LibreOffice."""
    from docx import Document

    program = find_libreoffice()
    if program is None:
        raise SystemExit("LibreOffice is needed to save the notes as a PDF.")
    notes = Document()
    for number, (level, heading, paragraphs) in enumerate(NOTES):
        # A section to a page, as lecture notes are often laid out, so a
        # retrieval question's answer is on a page of its own.
        if number > 1:
            notes.add_page_break()
        notes.add_heading(heading, level=0 if level == 1 else 1)
        for text in paragraphs:
            notes.add_paragraph(text)
    with tempfile.TemporaryDirectory() as work:
        source = Path(work) / f"{path.stem}.docx"
        notes.save(source)
        subprocess.run(
            [program, "--headless", "--convert-to", "pdf", "--outdir", work, str(source)],
            check=True,
            capture_output=True,
            timeout=120,
        )
        path.write_bytes((Path(work) / f"{path.stem}.pdf").read_bytes())


def main() -> None:
    build_cnn_deck(SAMPLES / "cnn_lecture.pptx")
    build_notes_pdf(SAMPLES / "linear_regression_notes.pdf")
    print(
        f"written to {SAMPLES / 'cnn_lecture.pptx'} and {SAMPLES / 'linear_regression_notes.pdf'}"
    )


if __name__ == "__main__":
    main()

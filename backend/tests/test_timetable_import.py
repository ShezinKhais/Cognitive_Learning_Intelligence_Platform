"""Unit tests for the timetable/roster parsing service.

Owner: Cyber 2. No DB, no auth — these exercise the parsing/conflict logic
in isolation, since app.services.timetable_import is deliberately DB-agnostic.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.core.errors import ValidationError
from app.services import timetable_import
from app.services.timetable_import import (
    detect_timetable_conflicts,
    match_names,
    parse_roster,
    parse_timetable,
)

TIMETABLE_HEADER = "course_code,lecturer,day,start_time,end_time,room\n"
ROSTER_HEADER = "student_email,student_name,course_code\n"


def _csv(header: str, *rows: str) -> bytes:
    return (header + "\n".join(rows)).encode("utf-8")


# --- parse_timetable ---------------------------------------------------


def test_parses_valid_timetable_csv():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. Orumchian,Monday,09:00,11:00,Room 12",
        "CSIT101,Dr. Smith,Tuesday,13:00,14:30,Room 4",
    )
    rows, rows_read = parse_timetable("timetable.csv", raw)
    assert rows_read == 2
    assert len(rows) == 2
    assert rows[0].course_code == "CSIT321"
    assert rows[0].lecturer == "Dr. Orumchian"
    assert rows[0].day == "Monday"


def test_rejects_non_csv_xlsx_extension():
    with pytest.raises(ValidationError):
        parse_timetable("timetable.pdf", b"whatever")


def test_rejects_empty_file():
    with pytest.raises(ValidationError):
        parse_timetable("timetable.csv", b"")


def test_rejects_missing_required_column():
    raw = b"course_code,lecturer,day,start_time,end_time\nCSIT321,X,Monday,09:00,11:00\n"
    with pytest.raises(ValidationError, match="missing required columns"):
        parse_timetable("timetable.csv", raw)


def test_normalises_short_day_name():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,X,Mon,09:00,11:00,Room 1",
    )

    rows, _ = parse_timetable(
        "timetable.csv",
        raw,
    )

    assert rows[0].day == "Monday"


def test_rejects_invalid_day():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,X,Funday,09:00,11:00,Room 1",
    )

    with pytest.raises(
        ValidationError,
        match="not a valid day",
    ):
        parse_timetable(
            "timetable.csv",
            raw,
        )


def test_rejects_bad_time_format():
    raw = _csv(TIMETABLE_HEADER, "CSIT321,X,Monday,9am,11am,Room 1")
    with pytest.raises(ValidationError, match="could not parse"):
        parse_timetable("timetable.csv", raw)


def test_rejects_end_before_start():
    raw = _csv(TIMETABLE_HEADER, "CSIT321,X,Monday,14:00,13:00,Room 1")
    with pytest.raises(ValidationError, match="not after"):
        parse_timetable("timetable.csv", raw)


def test_timetable_first_data_row_is_row_two():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,X,Monday,09:00,11:00,Room 1",
    )

    rows, _ = parse_timetable(
        "timetable.csv",
        raw,
    )

    assert rows[0].row_number == 2


def test_roster_first_data_row_is_row_two():
    raw = _csv(
        ROSTER_HEADER,
        "a.student@uni.test,A Student,CSIT321",
    )

    rows, _ = parse_roster(
        "roster.csv",
        raw,
    )

    assert rows[0].row_number == 2


def test_rejects_blank_required_field():
    raw = _csv(TIMETABLE_HEADER, ",X,Monday,09:00,11:00,Room 1")
    with pytest.raises(ValidationError):
        parse_timetable("timetable.csv", raw)


def test_parses_valid_timetable_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["course_code", "lecturer", "day", "start_time", "end_time", "room"])
    ws.append(["CSIT321", "Dr. Orumchian", "Monday", "09:00", "11:00", "Room 12"])
    buf = io.BytesIO()
    wb.save(buf)

    rows, rows_read = parse_timetable("timetable.xlsx", buf.getvalue())
    assert rows_read == 1
    assert rows[0].course_code == "CSIT321"


def test_xlsx_workbook_is_closed(monkeypatch):
    closed = False

    class FakeWorksheet:
        def iter_rows(self, values_only=True):
            return iter(
                [
                    [
                        "course_code",
                        "lecturer",
                        "day",
                        "start_time",
                        "end_time",
                        "room",
                    ],
                    [
                        "CSIT321",
                        "Dr. Orumchian",
                        "Monday",
                        "09:00",
                        "11:00",
                        "Room 12",
                    ],
                ]
            )

    class FakeWorkbook:
        active = FakeWorksheet()

        def close(self):
            nonlocal closed
            closed = True

    monkeypatch.setattr(
        timetable_import.openpyxl,
        "load_workbook",
        lambda *args, **kwargs: FakeWorkbook(),
    )

    parse_timetable(
        "timetable.xlsx",
        b"fake-xlsx-data",
    )

    assert closed is True


def test_rejects_xlsx_blank_header():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.append(
        [
            "course_code",
            "lecturer",
            "day",
            "start_time",
            "end_time",
            None,
        ]
    )

    ws.append(
        [
            "CSIT321",
            "Dr. A",
            "Monday",
            "09:00",
            "11:00",
            "Room 1",
        ]
    )

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    with pytest.raises(
        ValidationError,
        match="blank column",
    ):
        parse_timetable(
            "timetable.xlsx",
            buf.getvalue(),
        )


def test_rejects_xlsx_duplicate_header():
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.append(
        [
            "course_code",
            "lecturer",
            "day",
            "start_time",
            "end_time",
            "course_code",
        ]
    )

    ws.append(
        [
            "CSIT321",
            "Dr. A",
            "Monday",
            "09:00",
            "11:00",
            "CSIT999",
        ]
    )

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    with pytest.raises(
        ValidationError,
        match="duplicate column",
    ):
        parse_timetable(
            "timetable.xlsx",
            buf.getvalue(),
        )


def test_xlsx_skips_blank_trailing_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["course_code", "lecturer", "day", "start_time", "end_time", "room"])
    ws.append(["CSIT321", "Dr. Orumchian", "Monday", "09:00", "11:00", "Room 12"])
    ws.append([None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)

    rows, rows_read = parse_timetable("timetable.xlsx", buf.getvalue())
    assert rows_read == 1


# --- parse_roster --------------------------------------------------------


def test_parses_valid_roster_csv():
    raw = _csv(ROSTER_HEADER, "a.student@uni.test,A Student,CSIT321")
    rows, rows_read = parse_roster("roster.csv", raw)
    assert rows_read == 1
    assert rows[0].student_email == "a.student@uni.test"


@pytest.mark.parametrize(
    "email",
    [
        "not-an-email",
        "student@",
        "@uni.test",
        "student@@uni.test",
    ],
)
def test_rejects_invalid_email(email):
    raw = _csv(
        ROSTER_HEADER,
        f"{email},A Student,CSIT321",
    )

    with pytest.raises(
        ValidationError,
        match="valid email",
    ):
        parse_roster(
            "roster.csv",
            raw,
        )


# --- detect_timetable_conflicts ------------------------------------------


def test_detects_lecturer_double_booking():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. Orumchian,Monday,09:00,11:00,Room 12",
        "CSIT101,Dr. Orumchian,Monday,10:00,12:00,Room 4",
    )
    rows, _ = parse_timetable("timetable.csv", raw)
    conflicts = detect_timetable_conflicts(rows)
    assert len(conflicts) == 1
    assert "Dr. Orumchian" in conflicts[0]


def test_detects_room_double_booking():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. A,Monday,09:00,11:00,Room 12",
        "CSIT101,Dr. B,Monday,10:00,12:00,Room 12",
    )
    rows, _ = parse_timetable("timetable.csv", raw)
    conflicts = detect_timetable_conflicts(rows)
    assert len(conflicts) == 1
    assert "Room 'Room 12'" in conflicts[0]


def test_detects_multiple_overlaps_in_same_group():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. A,Monday,09:00,13:00,Room 1",
        "CSIT101,Dr. A,Monday,10:00,11:00,Room 2",
        "CSIT202,Dr. A,Monday,12:00,14:00,Room 3",
    )

    rows, _ = parse_timetable(
        "timetable.csv",
        raw,
    )

    conflicts = detect_timetable_conflicts(rows)

    lecturer_conflicts = [conflict for conflict in conflicts if conflict.startswith("Lecturer")]

    assert len(lecturer_conflicts) == 2


def test_no_conflict_for_back_to_back_sessions():
    """11:00-12:00 followed by 12:00-13:00 does not overlap."""
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. A,Monday,11:00,12:00,Room 12",
        "CSIT101,Dr. A,Monday,12:00,13:00,Room 12",
    )
    rows, _ = parse_timetable("timetable.csv", raw)
    assert detect_timetable_conflicts(rows) == []


def test_no_conflict_across_different_days():
    raw = _csv(
        TIMETABLE_HEADER,
        "CSIT321,Dr. A,Monday,09:00,11:00,Room 12",
        "CSIT101,Dr. A,Tuesday,09:00,11:00,Room 12",
    )
    rows, _ = parse_timetable("timetable.csv", raw)
    assert detect_timetable_conflicts(rows) == []


# --- match_names -----------------------------------------------------------


def test_match_names_exact_match():
    unmatched, matched = match_names(["Dr. Orumchian"], ["Dr. Orumchian"])
    assert unmatched == []
    assert matched == {"Dr. Orumchian": "Dr. Orumchian"}


def test_match_names_near_match_typo():
    unmatched, matched = match_names(["Dr Orumchian"], ["Dr. Orumchian"])
    assert unmatched == []
    assert matched["Dr Orumchian"] == "Dr. Orumchian"


def test_match_names_no_known_list_means_all_unmatched():
    unmatched, matched = match_names(["Dr. Orumchian", "Dr. Orumchian"], [])
    assert unmatched == ["Dr. Orumchian"]  # de-duplicated
    assert matched == {}


def test_match_names_below_threshold_is_unmatched():
    unmatched, matched = match_names(["Totally Different Name"], ["Dr. Orumchian"])
    assert unmatched == ["Totally Different Name"]
    assert matched == {}
